"""Excel sheet of the paper bets, kept to the basics. Sheet "Opklade": the current state on top (bank, profit, ROI,
bets, won-lost, pending), then one row per bet: kick-off (Belgrade time), tournament, who we back, opponent, odds,
stake, result, profit, bookmaker. Sheet "Po kategoriji": the same totals per Flashscore category. Sheet "Sve kvote": every book's latest
price for our side (px_<book> in the ledger), the best in green.
Stake = quarter Kelly of the bank at the moment the bet was logged = starting bank + profit of every bet already settled
then (settled = when the loop first saw the result), so a settled bet changes only the stakes of later bets. All cells
are values (computed here), so every viewer shows them.

usage: ledger_xlsx.py LEDGER_CSV OUT_XLSX [BANK=100000]   (dinars)
The same rows also go to OUT_XLSX with .json in place of .xlsx, the data of the online dashboard.
"""
import json
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from fs_common import player_name

SRC, OUT = sys.argv[1], sys.argv[2]
BANK = float(sys.argv[3]) if len(sys.argv) > 3 else 100000.0
KELLY = 0.25
SEG = {"ITF WOMEN - SINGLES": "Ženski ITF", "CHALLENGER MEN - SINGLES": "Muški Challenger", "ITF MEN - SINGLES": "Muški ITF",
       "CHALLENGER WOMEN - SINGLES": "WTA 125", "WTA - SINGLES": "WTA Tour", "ATP - SINGLES": "ATP Tour"}

BOOK_NAMES = {"bet365": "bet365", "Mozzartbetrs": "MozzartBet", "Admiralbet.rs": "AdmiralBet", "SoccerBet": "SoccerBet",
              "Superbet.rs": "Superbet", "1xBet.rs": "1xBet"}

L = pd.read_csv(SRC)
L["thr"] = L.rule.str.extract(r"EV>(\d+)%").astype(float)[0]
L = L.sort_values(["thr", "logged_at"]).drop_duplicates(["fs_mid", "side"])
L = L[L.tournament.str.split(":").str[0].str.strip().isin(SEG)].copy()
now = pd.Timestamp.now("UTC").tz_localize(None)
bg = lambda t: pd.Timestamp(t).tz_localize("UTC").tz_convert("Europe/Belgrade").tz_localize(None).to_pydatetime()

L["res"] = [{"void": "Povraćaj"}.get(str(r), "Dobijeno" if r == s else ("Izgubljeno" if str(r) in ("home", "away") else ""))
            for r, s in zip(L.result, L.side)]
L["res"] = [r or ("U toku" if pd.Timestamp(s) <= now else "Čeka") for r, s in zip(L.res, L.start)]
# settled before settlement times were recorded: kick-off + 3 h
L["settled"] = [pd.Timestamp(a) if isinstance(a, str) and a else pd.Timestamp(st) + pd.Timedelta(hours=3)
                for a, st in zip(L.get("settled_at", pd.Series(np.nan, index=L.index)), L.start)]
L["logged"] = pd.to_datetime(L.logged_at)
# stakes in the order the bets were logged: a bet settled before a later bet was logged is in that bet's bank
stake, profit = {}, {}
for i, b in L.sort_values("logged").iterrows():
    bank = BANK + sum(v for j, v in profit.items() if L.at[j, "settled"] <= b.logged)
    p, o = float(b.q_blend), float(b.odds)
    stake[i] = round(bank * KELLY * max(0.0, (p * o - 1) / (o - 1)))
    if b.res in ("Dobijeno", "Izgubljeno", "Povraćaj"):
        profit[i] = round(stake[i] * (o - 1)) if b.res == "Dobijeno" else (-stake[i] if b.res == "Izgubljeno" else 0)
L["stake"], L["profit"] = pd.Series(stake), pd.Series(profit)
L["cat"] = L.tournament.str.split(":").str[0].str.strip().map(SEG)
L = L.sort_values(["start", "logged"])

NAVY, INK, MUTED, LINE, BAND = "1B2A41", "1F2933", "7B8794", "E4E7EB", "F7F9FB"
GREEN, GREEN_BG, RED, RED_BG, AMBER, AMBER_BG, BLUE_BG = "1E7B34", "E3F4E8", "B42318", "FDE8E7", "8A5A00", "FFF4D6", "E8EEF8"
font = lambda **k: Font(**{"name": "Calibri", "size": 10, "color": INK, **k})
fill = lambda c: PatternFill("solid", fgColor=c)
line = Border(bottom=Side(style="thin", color=LINE))
STATUS = {"Dobijeno": (GREEN, GREEN_BG), "Izgubljeno": (RED, RED_BG), "Čeka": (AMBER, AMBER_BG), "U toku": (AMBER, AMBER_BG),
          "Povraćaj": (MUTED, BAND)}


def totals(g):
    s = g[g.res.isin(["Dobijeno", "Izgubljeno"])]
    staked = s.stake.sum()
    return {"profit": g.profit.sum(), "roi": g.profit.sum() / staked if staked else None, "n": len(g),
            "won": int((g.res == "Dobijeno").sum()), "lost": int((g.res == "Izgubljeno").sum()),
            "open": int(g.res.isin(["Čeka", "U toku"]).sum())}


def header(ws, row, names, widths, right):
    for j, (c, w) in enumerate(zip(names, widths), 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font, cell.fill = font(bold=True, color="FFFFFF"), fill(NAVY)
        cell.alignment = Alignment(horizontal="right" if c in right else "left", vertical="center", indent=1)
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[row].height = 22


wb = Workbook()
ws = wb.active
ws.title = "Opklade"
ws.sheet_view.showGridLines = False
T = totals(L)
ws["A1"] = "Opklade"
ws["A1"].font = font(size=16, bold=True, color=NAVY)
ws["A1"].alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 28
pos = lambda v: GREEN if v and v > 0 else (RED if v and v < 0 else INK)
cards = [("Banka", BANK + T["profit"], "#,##0", pos(T["profit"]), BLUE_BG),
         ("Dobit", T["profit"], "+#,##0;-#,##0;0", pos(T["profit"]), GREEN_BG if T["profit"] > 0 else RED_BG if T["profit"] < 0 else BAND),
         ("ROI", T["roi"], "+0.0%;-0.0%;0.0%", pos(T["roi"]), BAND),
         ("Opklade", T["n"], "0", INK, BAND), ("Dobijeno – izgubljeno", f'{T["won"]} – {T["lost"]}', "@", INK, BAND),
         ("Čeka", T["open"], "0", AMBER, AMBER_BG)]
# cards: label in row 3, value in row 4, over the columns of the bet table (A-I)
for (lab, v, fmt, color, bgc), (c0, c1) in zip(cards, [(1, 2), (3, 3), (4, 4), (5, 6), (7, 8), (9, 9)]):
    for r in (3, 4):
        if c1 > c0:
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        for c in range(c0, c1 + 1):
            ws.cell(row=r, column=c).fill = fill(bgc)
    a, b = ws.cell(row=3, column=c0, value=lab), ws.cell(row=4, column=c0, value=v if v is not None else "–")
    a.font, b.font = font(size=9, color=MUTED), font(size=16, bold=True, color=color)
    b.number_format = fmt
    a.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[3].height, ws.row_dimensions[4].height = 18, 30
ws.cell(row=5, column=1, value=f"Početna banka {BANK:,.0f} · četvrtina Kelly-ja · stanje {bg(now):%d.%m. %H:%M}".replace(",", "."))
ws.cell(row=5, column=1).font = font(size=9, color=MUTED)

COLS = [("Početak", 13, "dd.mm. hh:mm"), ("Turnir", 34, "@"), ("Igramo na", 22, "@"), ("Protiv", 22, "@"), ("Kvota", 8, "0.00"),
        ("Ulog", 10, "#,##0"), ("Rezultat", 12, "@"), ("Dobit", 11, "+#,##0;-#,##0;0"), ("Kladionica", 14, "@")]
RIGHT = {"Kvota", "Ulog", "Dobit"}
r0 = 8
header(ws, r0 - 1, [c for c, _, _ in COLS], [w for _, w, _ in COLS], RIGHT)
for i, b in enumerate(L.itertuples()):
    r = r0 + i
    pick, opp = (b.home, b.away) if b.side == "home" else (b.away, b.home)
    tname = b.tournament.split(":", 1)[-1].strip()
    vals = [bg(b.start), tname, player_name({"WU": pick}, "h").title(), player_name({"WU": opp}, "h").title(), float(b.odds),
            int(b.stake), b.res, None if np.isnan(b.profit) else int(b.profit), BOOK_NAMES.get(b.book, b.book)]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=j, value=v)
        cell.number_format = COLS[j - 1][2]
        cell.font = font(bold=j == 3)
        cell.border = line
        cell.alignment = Alignment(horizontal="right" if COLS[j - 1][0] in RIGHT else "left", vertical="center", indent=1)
        if i % 2:
            cell.fill = fill(BAND)
    fg, bgc = STATUS.get(b.res, (INK, BAND))
    st = ws.cell(row=r, column=7)
    st.font, st.fill = font(bold=True, color=fg), fill(bgc)
    if not np.isnan(b.profit) and b.profit:
        ws.cell(row=r, column=8).font = font(bold=True, color=GREEN if b.profit > 0 else RED)
    ws.row_dimensions[r].height = 20
last = r0 + max(len(L), 1) - 1
ws.freeze_panes = ws.cell(row=r0, column=1)
ws.auto_filter.ref = f"A{r0 - 1}:I{last}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# per category
wc = wb.create_sheet("Po kategoriji")
wc.sheet_view.showGridLines = False
wc["A1"] = "Po kategoriji"
wc["A1"].font = font(size=16, bold=True, color=NAVY)
wc.row_dimensions[1].height = 28
names = ["Kategorija", "Opklade", "Dobijeno", "Izgubljeno", "Čeka", "Dobit", "ROI"]
header(wc, 3, names, [20, 10, 10, 11, 8, 12, 9], set(names[1:]))
rows = [(cat, totals(L[L.cat == cat])) for cat in SEG.values()] + [("Ukupno", T)]
for i, (cat, t) in enumerate(rows):
    r = 4 + i
    vals = [cat, t["n"], t["won"], t["lost"], t["open"], t["profit"], t["roi"]]
    for j, v in enumerate(vals, 1):
        cell = wc.cell(row=r, column=j, value=v if v is not None else "–")
        cell.number_format = {6: "+#,##0;-#,##0;0", 7: "+0.0%;-0.0%;0.0%"}.get(j, "0" if j > 1 else "@")
        cell.font = font(bold=cat == "Ukupno" or j == 1, color=pos(v) if j in (6, 7) and v else INK)
        cell.border = line
        cell.alignment = Alignment(horizontal="left" if j == 1 else "right", vertical="center", indent=1)
        if cat == "Ukupno":
            cell.fill = fill(BLUE_BG)
        elif i % 2:
            cell.fill = fill(BAND)
    wc.row_dimensions[r].height = 20
# every book's latest price for our side (before kick-off), the best one in green
wk = wb.create_sheet("Sve kvote")
wk.sheet_view.showGridLines = False
wk["A1"] = "Sve kvote"
wk["A1"].font = font(size=16, bold=True, color=NAVY)
wk.row_dimensions[1].height = 28
wk["A2"] = "zadnja kvota pre početka meča za našu stranu · zeleno = najbolja"
wk["A2"].font = font(size=9, color=MUTED)
books = [b for b in BOOK_NAMES if f"px_{b}" in L]
names = ["Početak", "Igramo na", "Protiv", "Uzeli smo"] + [BOOK_NAMES[b] for b in books]
header(wk, 4, names, [13, 22, 22, 17] + [11] * len(books), set(names[3:]))
for i, b in enumerate(L.itertuples()):
    r = 5 + i
    pick, opp = (b.home, b.away) if b.side == "home" else (b.away, b.home)
    px = [float(L.at[b.Index, f"px_{bk}"]) for bk in books]
    best = np.nanmax(px) if not all(np.isnan(px)) else None
    vals = [bg(b.start), player_name({"WU": pick}, "h").title(), player_name({"WU": opp}, "h").title(),
            f"{float(b.odds):.2f} {BOOK_NAMES.get(b.book, b.book)}"] + [None if np.isnan(v) else v for v in px]
    for j, v in enumerate(vals, 1):
        cell = wk.cell(row=r, column=j, value=v)
        cell.number_format = "dd.mm. hh:mm" if j == 1 else ("0.00" if j > 4 else "@")
        cell.font = font(bold=j == 2)
        cell.border = line
        cell.alignment = Alignment(horizontal="right" if j >= 4 else "left", vertical="center", indent=1)
        if i % 2:
            cell.fill = fill(BAND)
        if j > 4 and v is not None and best is not None and abs(v - best) < 1e-9:
            cell.font, cell.fill = font(bold=True, color=GREEN), fill(GREEN_BG)
    wk.row_dimensions[r].height = 20
wk.freeze_panes = wk.cell(row=5, column=1)
for sh in (wc, wk):
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
rows = []
for b in L.itertuples():
    pick, opp = (b.home, b.away) if b.side == "home" else (b.away, b.home)
    last = None if pd.isna(b.close) else float(b.close)
    rows.append({"id": b.fs_mid, "start": f"{bg(b.start):%Y-%m-%dT%H:%M}", "category": b.cat,
                 "tournament": b.tournament.split(":", 1)[-1].strip(),
                 "pick": player_name({"WU": pick}, "h").title(), "opponent": player_name({"WU": opp}, "h").title(),
                 "odds": float(b.odds), "book": BOOK_NAMES.get(b.book, b.book), "stake": int(b.stake),
                 "status": b.res, "profit": None if np.isnan(b.profit) else int(b.profit),
                 "p_model": round(float(b.q_blend), 4), "ev": round(float(b.ev), 4), "last_odds": last,
                 "clv": None if not last else round(float(b.odds) / last - 1, 4)})
with open(os.path.splitext(OUT)[0] + ".json", "w") as f:
    json.dump({"bank_start": BANK, "as_of": f"{bg(now):%Y-%m-%dT%H:%M}", "rows": rows}, f, ensure_ascii=False, indent=1)
print(f"{len(L)} bets -> {OUT}")
